"""A folha em arquivo: PDF e DOCX (ADR-078). Sem banco — `core.exportacao` só recebe um dicionário."""

import io
from types import SimpleNamespace

from django.test import SimpleTestCase
from docx import Document

from core import exportacao

FOLHA = {
    "titulo": "DECLARAÇÃO", "subtitulo": "",
    "blocos": [("", "Declaro, para fins de comprovação, que Marcos Tela compareceu a atendimento psicológico.")],
    "local_e_data": "Belo Horizonte, 18 de setembro de 2026",
    "assinatura_nome": "Ana Souza", "assinatura_crp": "CRP 04/12345", "tracos": False,
    "timbre": {"nome": "Clínica Exemplo", "endereco": "Rua das Flores, 10 · Centro · Belo Horizonte/MG"},
}


def texto_do_docx(conteudo: bytes) -> str:
    documento = Document(io.BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    for tabela in documento.tables:
        partes += [celula.text for linha in tabela.rows for celula in linha.cells]
    for secao in documento.sections:
        partes += [p.text for p in secao.header.paragraphs]
    return "\n".join(partes)


class Formatos(SimpleTestCase):
    def test_so_pdf_e_docx(self):
        self.assertEqual(set(exportacao.GERADORES), {"pdf", "docx"})

    def test_pdf_e_um_pdf(self):
        conteudo = exportacao.gerar_pdf(FOLHA)
        self.assertTrue(conteudo.startswith(b"%PDF"))
        self.assertGreater(len(conteudo), 1000)

    def test_docx_traz_timbre_titulo_texto_data_e_assinatura(self):
        texto = texto_do_docx(exportacao.gerar_docx(FOLHA))
        for trecho in ["Clínica Exemplo", "DECLARAÇÃO", "compareceu a atendimento", "18 de setembro de 2026",
                       "Ana Souza", "CRP 04/12345"]:
            with self.subTest(trecho=trecho):
                self.assertIn(trecho, texto)

    def test_documento_itemizado_leva_o_titulo_de_cada_bloco(self):
        folha = {**FOLHA, "titulo": "LAUDO PSICOLÓGICO",
                 "blocos": [("Identificação", "Nome: Marcos Tela"), ("Análise", "Texto da análise.")]}
        texto = texto_do_docx(exportacao.gerar_docx(folha))
        self.assertIn("ANÁLISE", texto)
        self.assertIn("Texto da análise.", texto)
        self.assertTrue(exportacao.gerar_pdf(folha).startswith(b"%PDF"))


class MarcasDaFolha(SimpleTestCase):
    def test_rascunho_sai_marcado_como_rascunho(self):
        self.assertIn("RASCUNHO", texto_do_docx(exportacao.gerar_docx({**FOLHA, "rascunho": True})))
        self.assertNotIn("RASCUNHO", texto_do_docx(exportacao.gerar_docx(FOLHA)))
        self.assertTrue(exportacao.gerar_pdf({**FOLHA, "rascunho": True}).startswith(b"%PDF"))

    def test_atestado_fecha_o_espaco_com_tracos(self):
        self.assertIn("----", texto_do_docx(exportacao.gerar_docx({**FOLHA, "tracos": True})))
        self.assertNotIn("----", texto_do_docx(exportacao.gerar_docx(FOLHA)))

    def test_segunda_assinatura_vem_junto(self):
        folha = {**FOLHA, "segunda_assinatura": "Assinatura de quem recebeu o documento"}
        self.assertIn("Assinatura de quem recebeu o documento", texto_do_docx(exportacao.gerar_docx(folha)))

    def test_folha_sem_timbre_e_sem_data_nao_quebra(self):
        folha = {"titulo": "PARECER", "blocos": [("", "Texto.")], "assinatura_nome": "Ana", "assinatura_crp": ""}
        self.assertTrue(exportacao.gerar_pdf(folha).startswith(b"%PDF"))
        self.assertIn("PARECER", texto_do_docx(exportacao.gerar_docx(folha)))

    def test_caractere_fora_do_latim_nao_derruba_o_pdf(self):
        folha = {**FOLHA, "blocos": [("", "Seta → e emoji 🙂 no meio do texto.")]}
        self.assertTrue(exportacao.gerar_pdf(folha).startswith(b"%PDF"))


class Timbre(SimpleTestCase):
    def test_timbre_vem_do_perfil(self):
        perfil = SimpleNamespace(nome_clinica="Clínica Exemplo", logradouro="Rua das Flores", numero="10",
                                 complemento="sala 2", bairro="Centro", cidade="Belo Horizonte", uf="MG")
        self.assertEqual(exportacao.timbre_de(perfil), {
            "nome": "Clínica Exemplo", "endereco": "Rua das Flores, 10 — sala 2 · Centro · Belo Horizonte/MG"})

    def test_perfil_sem_clinica_nao_tem_timbre(self):
        self.assertEqual(exportacao.timbre_de(SimpleNamespace()), {"nome": "", "endereco": ""})
