"""Da folha composta ao arquivo: PDF e DOCX (ADR-078).

Mora em `core` porque serve a mais de um app — documentos psicológicos e prontuário — e não conhece nenhum: recebe
uma **folha**, que é um dicionário simples, e devolve bytes. Quem compõe a folha é o app dono do conteúdo.

A folha::

    {"titulo": "DECLARAÇÃO", "subtitulo": "", "blocos": [("", "Declara-se…"), ("Análise", "…")],
     "local_e_data": "Belo Horizonte, 31 de maio de 2025",
     "assinatura_nome": "Ana Ribeiro", "assinatura_crp": "CRP 04/123456", "segunda_assinatura": "",
     "tracos": False, "timbre": {"nome": "…", "endereco": "…"}, "rascunho": False}

**Forma, conforme a Res. CFP nº 06/2019:** título centralizado; texto justificado; no atestado, o espaço que sobra é
fechado com traços (Art. 10, §5º); encerramento com local, data, nome e CRP; **laudas numeradas** no rodapé
("1/3") — a rubrica e o carimbo são de punho, ou a assinatura é eletrônica. Rascunho sai marcado como rascunho:
arquivo solto não pode se passar por documento emitido.

O PDF usa as fontes-padrão do formato (Times), que cobrem o português inteiro. Caractere fora delas — um emoji
colado no texto — é trocado por "?" em vez de derrubar a geração.
"""

import io
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TIPO_PDF = "application/pdf"
TIPO_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
# Os traços que fecham o atestado, em grupos: um bloco único de 90 traços não cabe no resto da linha e o
# editor o joga inteiro para a linha de baixo, deixando um buraco justamente onde os traços deviam estar.
_TRACOS = " ".join(["----"] * 28)
_AVISO_DE_RASCUNHO = "RASCUNHO — documento ainda não emitido"


def timbre_de(usuario) -> dict:
    """O cabeçalho da folha, a partir do perfil: nome da clínica e endereço. Vazio, não há timbre."""
    nome = getattr(usuario, "nome_clinica", "") or ""
    rua = getattr(usuario, "logradouro", "") or ""
    partes = []
    if rua:
        numero = getattr(usuario, "numero", "")
        complemento = getattr(usuario, "complemento", "")
        partes.append(rua + (f", {numero}" if numero else "") + (f" — {complemento}" if complemento else ""))
        if getattr(usuario, "bairro", ""):
            partes.append(usuario.bairro)
        if getattr(usuario, "cidade", ""):
            partes.append(usuario.cidade + (f"/{usuario.uf}" if getattr(usuario, "uf", "") else ""))
    return {"nome": nome, "endereco": " · ".join(partes)}


def _assinaturas(folha: dict) -> list:
    """Cada assinatura é a lista de linhas que vai embaixo do traço. A de quem recebe ou autoriza vem primeiro."""
    assinaturas = []
    if folha.get("segunda_assinatura"):
        assinaturas.append([folha["segunda_assinatura"]])
    assinaturas.append([linha for linha in [folha.get("assinatura_nome", ""), folha.get("assinatura_crp", "")] if linha])
    return assinaturas


def _paragrafos(texto: str) -> list:
    return [paragrafo for paragrafo in (texto or "").split("\n\n") if paragrafo.strip()]


# --- PDF ---------------------------------------------------------------------------------------------------------------

def _latim(texto: str) -> str:
    """As fontes-padrão do PDF vão até o Windows-1252. O que passar disso vira "?", em vez de erro."""
    return (texto or "").encode("cp1252", "replace").decode("cp1252")


def _marcado(texto: str) -> str:
    return escape(_latim(texto)).replace("\n", "<br/>")


class _LaudasNumeradas(canvas.Canvas):
    """Numera "1/3" no rodapé — o total só se sabe no fim, então as páginas são guardadas e fechadas juntas."""

    rascunho = False

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._laudas = []

    def showPage(self):
        self._laudas.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._laudas)
        for estado in self._laudas:
            self.__dict__.update(estado)
            self._fechar_lauda(total)
            super().showPage()
        super().save()

    def _fechar_lauda(self, total: int) -> None:
        largura, altura = A4
        self.setFont("Times-Roman", 9)
        self.drawCentredString(largura / 2, 1.2 * cm, f"{self._pageNumber}/{total}")
        if self.rascunho:
            self.saveState()
            self.setFont("Helvetica-Bold", 84)
            self.setFillColorRGB(0.78, 0.16, 0.16, alpha=0.12)
            self.translate(largura / 2, altura / 2)
            self.rotate(30)
            self.drawCentredString(0, 0, "RASCUNHO")
            self.restoreState()


def gerar_pdf(folha: dict) -> bytes:
    base = ParagraphStyle("base", fontName="Times-Roman", fontSize=12, leading=19)
    estilos = {
        "timbre_nome": ParagraphStyle("timbre_nome", parent=base, fontName="Times-Bold", alignment=TA_CENTER, leading=15),
        "timbre": ParagraphStyle("timbre", parent=base, fontSize=10, leading=13, alignment=TA_CENTER, textColor="#444444"),
        "titulo": ParagraphStyle("titulo", parent=base, fontName="Times-Bold", fontSize=13, alignment=TA_CENTER),
        "subtitulo": ParagraphStyle("subtitulo", parent=base, fontName="Times-Italic", alignment=TA_CENTER),
        "item": ParagraphStyle("item", parent=base, fontName="Times-Bold", fontSize=11.5, spaceBefore=14, spaceAfter=2),
        "texto": ParagraphStyle("texto", parent=base, alignment=TA_JUSTIFY, spaceAfter=8),
        "local": ParagraphStyle("local", parent=base, alignment=TA_RIGHT, spaceBefore=26),
        "assinatura": ParagraphStyle("assinatura", parent=base, alignment=TA_CENTER, leading=15),
    }
    historia = []

    timbre = folha.get("timbre") or {}
    if timbre.get("nome") or timbre.get("endereco"):
        if timbre.get("nome"):
            historia.append(Paragraph(_marcado(timbre["nome"]), estilos["timbre_nome"]))
        if timbre.get("endereco"):
            historia.append(Paragraph(_marcado(timbre["endereco"]), estilos["timbre"]))
        historia += [Spacer(1, 4), HRFlowable(width="100%", thickness=0.6, color="#999999"), Spacer(1, 18)]

    historia.append(Paragraph(_marcado(folha["titulo"]), estilos["titulo"]))
    if folha.get("subtitulo"):
        historia.append(Paragraph(_marcado(folha["subtitulo"]), estilos["subtitulo"]))
    historia.append(Spacer(1, 16))

    blocos = folha.get("blocos", [])
    for posicao, (titulo, texto) in enumerate(blocos):
        if titulo:
            historia.append(Paragraph(_marcado(titulo.upper()), estilos["item"]))
            historia.append(HRFlowable(width="100%", thickness=0.4, color="#bbbbbb", spaceAfter=4))
        paragrafos = _paragrafos(texto)
        for indice, paragrafo in enumerate(paragrafos):
            ultimo = posicao == len(blocos) - 1 and indice == len(paragrafos) - 1
            fecho = f" {_TRACOS}" if folha.get("tracos") and ultimo else ""
            historia.append(Paragraph(_marcado(paragrafo) + fecho, estilos["texto"]))

    fecho = []
    if folha.get("local_e_data"):
        fecho.append(Paragraph(_marcado(folha["local_e_data"]) + ".", estilos["local"]))
    fecho.append(Spacer(1, 46))
    # Cada assinatura tem o traço do tamanho de uma assinatura — 7,5 cm —, e não da largura da folha. Entre duas,
    # uma coluna vazia separa os traços.
    celulas, larguras, tracos = [], [], []
    for linhas in _assinaturas(folha):
        if celulas:
            celulas.append("")
            larguras.append(1.2 * cm)
        tracos.append(len(celulas))
        celulas.append(Paragraph("<br/>".join(_marcado(linha) for linha in linhas), estilos["assinatura"]))
        larguras.append(7.5 * cm)
    tabela = Table([celulas], colWidths=larguras, hAlign="CENTER")
    tabela.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, 0), 4)]
                               + [("LINEABOVE", (coluna, 0), (coluna, 0), 0.7, "#111111") for coluna in tracos]))
    fecho.append(tabela)
    # A assinatura não fica sozinha numa lauda: o último parágrafo do texto vai junto com ela. Folha só com
    # assinatura é o que permite trocar o conteúdo que veio antes.
    if historia and isinstance(historia[-1], Paragraph):
        fecho.insert(0, historia.pop())
        # Se esse parágrafo era o único do item, o título do item vem junto — título não fica órfão no pé da lauda.
        if len(historia) >= 2 and isinstance(historia[-1], HRFlowable):
            fecho[0:0] = [historia[-2], historia[-1]]
            del historia[-2:]
    historia.append(KeepTogether(fecho))

    saida = io.BytesIO()
    documento = SimpleDocTemplate(saida, pagesize=A4, leftMargin=2.5 * cm, rightMargin=2.5 * cm, topMargin=2.5 * cm,
                                  bottomMargin=2.3 * cm, title=_latim(folha["titulo"]))
    tela = type("Laudas", (_LaudasNumeradas,), {"rascunho": bool(folha.get("rascunho"))})
    documento.build(historia, canvasmaker=tela)
    return saida.getvalue()


# --- DOCX --------------------------------------------------------------------------------------------------------------

def _linha_embaixo(paragrafo, cor: str = "999999") -> None:
    borda = OxmlElement("w:pBdr")
    linha = OxmlElement("w:bottom")
    for atributo, valor in (("w:val", "single"), ("w:sz", "6"), ("w:space", "2"), ("w:color", cor)):
        linha.set(qn(atributo), valor)
    borda.append(linha)
    paragrafo._p.get_or_add_pPr().append(borda)


def _campo_do_word(paragrafo, instrucao: str) -> None:
    """Um campo calculado pelo Word — `PAGE`, `NUMPAGES` —, que é como o rodapé sabe o número da lauda."""
    for tipo, texto in (("begin", None), (None, instrucao), ("end", None)):
        corrida = paragrafo.add_run()
        if tipo:
            marca = OxmlElement("w:fldChar")
            marca.set(qn("w:fldCharType"), tipo)
        else:
            marca = OxmlElement("w:instrText")
            marca.set(qn("xml:space"), "preserve")
            marca.text = texto
        corrida._r.append(marca)


def _escrever(paragrafo, texto: str, **fonte):
    """O texto com as quebras de linha dele — no Word, quebra dentro do parágrafo, não parágrafo novo."""
    corrida = None
    for indice, linha in enumerate(texto.split("\n")):
        corrida = paragrafo.add_run(linha)
        for atributo, valor in fonte.items():
            setattr(corrida.font, atributo, valor)
        if indice < texto.count("\n"):
            corrida.add_break()
    return corrida


def gerar_docx(folha: dict) -> bytes:
    documento = Document()
    secao = documento.sections[0]
    secao.page_width, secao.page_height = Cm(21), Cm(29.7)
    secao.left_margin = secao.right_margin = secao.top_margin = Cm(2.5)
    secao.bottom_margin = Cm(2.3)

    normal = documento.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _campo_do_word(rodape, "PAGE")
    rodape.add_run("/")
    _campo_do_word(rodape, "NUMPAGES")
    if folha.get("rascunho"):
        aviso = secao.header.paragraphs[0]
        aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escrever(aviso, _AVISO_DE_RASCUNHO, bold=True, size=Pt(10)).font.color.rgb = RGBColor(0xB0, 0x28, 0x28)

    timbre = folha.get("timbre") or {}
    if timbre.get("nome") or timbre.get("endereco"):
        cabecalho = documento.add_paragraph()
        cabecalho.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cabecalho.paragraph_format.line_spacing = 1.1
        if timbre.get("nome"):
            _escrever(cabecalho, timbre["nome"], bold=True)
        if timbre.get("endereco"):
            if timbre.get("nome"):
                cabecalho.add_run().add_break()
            _escrever(cabecalho, timbre["endereco"], size=Pt(10))
        _linha_embaixo(cabecalho)
        cabecalho.paragraph_format.space_after = Pt(18)

    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _escrever(titulo, folha["titulo"], bold=True, size=Pt(13))
    if folha.get("subtitulo"):
        subtitulo = documento.add_paragraph()
        subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escrever(subtitulo, folha["subtitulo"], italic=True)

    blocos = folha.get("blocos", [])
    for posicao, (nome, texto) in enumerate(blocos):
        if nome:
            item = documento.add_paragraph()
            item.paragraph_format.space_before = Pt(14)
            item.paragraph_format.keep_with_next = True
            _escrever(item, nome.upper(), bold=True, size=Pt(11.5))
            _linha_embaixo(item, "BBBBBB")
        paragrafos = _paragrafos(texto)
        for indice, trecho in enumerate(paragrafos):
            paragrafo = documento.add_paragraph()
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            ultimo = posicao == len(blocos) - 1 and indice == len(paragrafos) - 1
            _escrever(paragrafo, trecho + (f" {_TRACOS}" if folha.get("tracos") and ultimo else ""))

    if folha.get("local_e_data"):
        local = documento.add_paragraph()
        local.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        local.paragraph_format.space_before = Pt(26)
        local.paragraph_format.keep_with_next = True
        _escrever(local, folha["local_e_data"] + ".")

    assinaturas = _assinaturas(folha)
    tabela = documento.add_table(rows=1, cols=len(assinaturas))
    tabela.autofit = True
    for celula, linhas in zip(tabela.rows[0].cells, assinaturas):
        paragrafo = celula.paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragrafo.paragraph_format.space_before = Pt(44)
        paragrafo.paragraph_format.line_spacing = 1.15
        _escrever(paragrafo, "_" * 38 + "\n" + "\n".join(linhas))

    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


GERADORES = {"pdf": (gerar_pdf, TIPO_PDF), "docx": (gerar_docx, TIPO_DOCX)}
